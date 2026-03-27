from __future__ import annotations

import csv
import re
from pathlib import Path


DEFAULT_PRODUCER_LIBRARY_ROWS: tuple[dict[str, str], ...] = (
    {"kod": "EG-U999-ST2-18", "typ": "plyta", "nazwa": "EGGER U999 ST2", "producent": "EGGER", "parametry": "U999 ST2", "grubosc": "18", "cena_zl": "289"},
    {"kod": "EG-H3303-ST10-18", "typ": "front", "nazwa": "EGGER H3303 ST10", "producent": "EGGER", "parametry": "H3303 ST10", "grubosc": "18", "cena_zl": "319"},
    {"kod": "SK-D2020-PE-18", "typ": "plyta", "nazwa": "Swisskrono D2020 PE", "producent": "Swisskrono", "parametry": "D2020 PE", "grubosc": "18", "cena_zl": "249"},
    {"kod": "SK-U190-PE-18", "typ": "front", "nazwa": "Swisskrono U190 PE", "producent": "Swisskrono", "parametry": "U190 PE", "grubosc": "18", "cena_zl": "269"},
    {"kod": "BL-TANDEMBOX-M", "typ": "okucie", "nazwa": "BLUM TANDEMBOX M", "producent": "BLUM", "parametry": "prowadnica + bok", "grubosc": "", "cena_zl": "189"},
    {"kod": "PK-MAGIC-CORNER", "typ": "okucie", "nazwa": "PEKA Magic Corner", "producent": "PEKA", "parametry": "narozny", "grubosc": "", "cena_zl": "1599"},
    {"kod": "KB-LEMANS", "typ": "okucie", "nazwa": "Kessebohmer LeMans", "producent": "Kessebohmer", "parametry": "narozny", "grubosc": "", "cena_zl": "1499"},
    {"kod": "SV-LOOX-16", "typ": "system_przesuwny", "nazwa": "SEVROLL Loox 16", "producent": "SEVROLL", "parametry": "tor + rolki", "grubosc": "", "cena_zl": "379"},
    {"kod": "WR-CONFIRMAT-7X50", "typ": "okucie", "nazwa": "WURTH Confirmat 7x50", "producent": "WURTH", "parametry": "wkrety", "grubosc": "", "cena_zl": "59"},
    {"kod": "AD-PUR-ANTISCRATCH", "typ": "lakier", "nazwa": "ADLER PUR-Antiscratch", "producent": "ADLER", "parametry": "RAL 9010", "grubosc": "", "cena_zl": "329"},
)


class ProducerLibraryService:
    def normalize_rows(self, rows: list[dict[str, str]]) -> list[dict[str, str]]:
        out: list[dict[str, str]] = []
        for entry in rows:
            if not isinstance(entry, dict):
                continue
            out.append(
                {
                    "kod": str(entry.get("kod", "") or "").strip(),
                    "typ": str(entry.get("typ", "") or "").strip().lower(),
                    "nazwa": str(entry.get("nazwa", "") or "").strip(),
                    "producent": str(entry.get("producent", "") or "").strip(),
                    "parametry": str(entry.get("parametry", "") or "").strip(),
                    "grubosc": str(entry.get("grubosc", "") or "").strip(),
                    "cena_zl": str(entry.get("cena_zl", "") or "").strip(),
                }
            )
        return out

    @staticmethod
    def normalize_header_name(value: str) -> str:
        raw = str(value or "").strip().lower()
        raw = raw.replace("ą", "a").replace("ć", "c").replace("ę", "e").replace("ł", "l")
        raw = raw.replace("ń", "n").replace("ó", "o").replace("ś", "s").replace("ż", "z").replace("ź", "z")
        raw = raw.replace(" ", "_")
        return re.sub(r"[^a-z0-9_]+", "", raw)

    def guess_import_mapping(self, headers: list[str]) -> dict[str, str]:
        normalized = {header: self.normalize_header_name(header) for header in headers}
        aliases: dict[str, tuple[str, ...]] = {
            "kod": ("kod", "code", "symbol", "sku", "indeks"),
            "typ": ("typ", "material_type", "rodzaj", "kategoria"),
            "nazwa": ("nazwa", "name", "nazwa_materialu", "material"),
            "producent": ("producent", "manufacturer", "brand", "marka"),
            "parametry": ("parametry", "parametr", "spec", "specyfikacja", "dekor"),
            "grubosc": ("grubosc", "thickness", "mm", "grub"),
            "cena_zl": ("cena", "cena_zl", "price", "price_zl", "cost"),
        }
        mapping: dict[str, str] = {}
        for target, names in aliases.items():
            for header, norm in normalized.items():
                if norm in names:
                    mapping[target] = header
                    break
            if target not in mapping:
                for header, norm in normalized.items():
                    if any(token in norm for token in names):
                        mapping[target] = header
                        break
        return mapping

    def load_rows_from_csv_or_tsv(self, file_path: Path) -> tuple[list[str], list[list[str]]]:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        sample = text[:4096]
        delimiter = ";"
        try:
            sniffed = csv.Sniffer().sniff(sample, delimiters=";,\t,")
            delimiter = str(sniffed.delimiter or ";")
        except Exception:
            if "\t" in sample and sample.count("\t") > sample.count(";"):
                delimiter = "\t"
            elif "," in sample and sample.count(",") > sample.count(";"):
                delimiter = ","
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        rows = [[str(cell or "").strip() for cell in row] for row in reader if any(str(cell or "").strip() for cell in row)]
        if not rows:
            return [], []
        return rows[0], rows[1:]

    def load_rows_from_xlsx(self, file_path: Path) -> tuple[list[str], list[list[str]]]:
        try:
            from openpyxl import load_workbook  # type: ignore
        except Exception as exc:
            raise RuntimeError("Brak biblioteki openpyxl do odczytu XLSX.") from exc
        wb = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        sheet = wb.active
        if sheet is None:
            return [], []
        rows_raw = list(sheet.iter_rows(values_only=True))
        rows = [[str(cell if cell is not None else "").strip() for cell in row] for row in rows_raw if any(cell is not None and str(cell).strip() for cell in row)]
        if not rows:
            return [], []
        return rows[0], rows[1:]

    def rows_from_import_buffer(self, headers: list[str], data_rows: list[list[str]], mapping: dict[str, str]) -> list[dict[str, str]]:
        if not headers or not data_rows:
            return []
        header_to_index = {header: idx for idx, header in enumerate(headers)}
        out: list[dict[str, str]] = []
        for row_values in data_rows:
            row: dict[str, str] = {}
            for target in ("kod", "typ", "nazwa", "producent", "parametry", "grubosc", "cena_zl"):
                source_header = str(mapping.get(target, "") or "")
                source_idx = header_to_index.get(source_header, -1)
                row[target] = str(row_values[source_idx] if 0 <= source_idx < len(row_values) else "").strip()
            row["typ"] = row["typ"].lower()
            if row["typ"] and row["nazwa"]:
                out.append(row)
        return out

    def merge_rows(self, existing_rows: list[dict[str, str]], incoming_rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], int]:
        merged = self.normalize_rows(existing_rows)
        added = 0
        known = {
            (
                str(entry.get("kod", "") or "").strip().lower(),
                str(entry.get("typ", "") or "").strip().lower(),
                str(entry.get("nazwa", "") or "").strip().lower(),
                str(entry.get("producent", "") or "").strip().lower(),
            )
            for entry in merged
        }
        for entry in self.normalize_rows(incoming_rows):
            key = (
                str(entry.get("kod", "") or "").strip().lower(),
                str(entry.get("typ", "") or "").strip().lower(),
                str(entry.get("nazwa", "") or "").strip().lower(),
                str(entry.get("producent", "") or "").strip().lower(),
            )
            if not key[1] or not key[2] or key in known:
                continue
            merged.append(entry)
            known.add(key)
            added += 1
        return merged, added

    def filter_rows(self, rows: list[dict[str, str]], producer: str, needle: str) -> list[dict[str, str]]:
        producer_norm = str(producer or "").strip().lower()
        needle_norm = str(needle or "").strip().lower()
        out: list[dict[str, str]] = []
        for row in self.normalize_rows(rows):
            if producer_norm and str(row.get("producent", "") or "").strip().lower() != producer_norm:
                continue
            if needle_norm:
                haystack = " | ".join((row["kod"], row["typ"], row["nazwa"], row["producent"], row["parametry"]))
                if needle_norm not in haystack.lower():
                    continue
            out.append(row)
        return out

    @staticmethod
    def extract_library_code(text: str) -> str:
        raw = str(text or "")
        match = re.search(r"\[KOD:([^\]]+)\]", raw)
        return str(match.group(1)).strip() if match else ""

    def compose_parametry_with_code(self, entry: dict[str, str]) -> str:
        base = str(entry.get("parametry", "") or "").strip()
        code = str(entry.get("kod", "") or "").strip()
        if not code:
            return base
        if f"[KOD:{code}]" in base:
            return base
        return (base + f" [KOD:{code}]").strip()

    def update_material_prices(self, material_rows: list[dict[str, str]], library_rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], int]:
        by_code = {
            str(entry.get("kod", "") or "").strip().lower(): entry
            for entry in self.normalize_rows(library_rows)
            if str(entry.get("kod", "") or "").strip()
        }
        by_triplet = {
            (
                str(entry.get("typ", "") or "").strip().lower(),
                str(entry.get("nazwa", "") or "").strip().lower(),
                str(entry.get("producent", "") or "").strip().lower(),
            ): entry
            for entry in self.normalize_rows(library_rows)
        }

        updated = 0
        out: list[dict[str, str]] = []
        for row in material_rows:
            item = dict(row)
            row_typ = str(item.get("typ", "") or "").strip().lower()
            row_name = str(item.get("nazwa", "") or "").strip().lower()
            row_prod = str(item.get("producent", "") or "").strip().lower()
            row_param = str(item.get("parametry", "") or "")

            source = None
            code = self.extract_library_code(row_param).lower()
            if code and code in by_code:
                source = by_code[code]
            else:
                source = by_triplet.get((row_typ, row_name, row_prod))

            price = str(source.get("cena_zl", "") or "").strip() if source else ""
            if price:
                if str(item.get("cena_zl", "") or "").strip() != price:
                    updated += 1
                item["cena_zl"] = price
            out.append(item)
        return out, updated
